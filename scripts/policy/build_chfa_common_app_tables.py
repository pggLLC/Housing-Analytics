#!/usr/bin/env python3
"""
scripts/policy/build_chfa_common_app_tables.py

Copies the reference tables CHFA builds into its Common Application workbook
into data/policy/chfa-common-app-tables.json, so the Deal Calculator and the
Finder can use the same numbers an applicant's workbook will.

Why the workbook: CHFA publishes these tables nowhere else. The QAP describes
Method Three (cost basis limits) and the two Round One location scores, but the
per-bedroom limits, the county zone factors, the county lowest-AMI points and
the city cost-burden tiers exist only on the workbook's hidden Admin tab.

What is copied (Admin tab of "2027 Common Application - Housing Tax Credit
Development - V9", released October 2026):
  - Basis_Limit_Test_Table: base limit per unit by bedrooms, elevator and
    walk-up (Lookup tab Building_Type_Table: Elevator = column 1, Walk-up = 2).
  - Zone_Factor_Table: the factor for zones 1-3.
  - The county table: each county's zone, its "Low income Score" (the
    lowest-AMI county points) and CHFA's DDA flag.
  - City_Table: the cost-burden score for each city, "Updated per new
    methodology 8/12/26".
How the workbook uses them (Calculations tab, rows 497-516 and 312, 316):
  limit per unit = TRUNC(zone factor x base limit for the bedroom count and
  building type); adjusted eligible basis = sum of units x limit per unit; then
  x 1.3 if QCT or DDA, x applicable fraction, x applicable percentage = the
  Method Three annual credit. The city score is looked up by the exact city
  name; a city not in the table scores 0 ("CITY MISSING").

The workbook follows the 2027-28 QAP Third Draft. The plan is not adopted, so
the output carries status "draft" and every consumer must say so.

Run by hand when CHFA posts a new workbook (chfa-qap-watch opens an issue):
  python3 scripts/policy/build_chfa_common_app_tables.py path/to/workbook.xlsx
The SHA-256 of the file read is recorded, so a reviewer can check it against
data/audit/chfa-qap-watch.json.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'data' / 'policy' / 'chfa-common-app-tables.json'
SOURCE_URL = 'https://www.chfainfo.com/getattachment/aaa81775-083b-401b-8fba-7474f13624d9/2027-Common-Application-for-Housing-Tax-Credit-Development-V9.xlsx'  # noqa: E501 (one line so the URL sweep reads it whole)
PAGE_URL = 'https://www.chfainfo.com/rental-housing/housing-credit/application'
BEDROOMS = ['0 bedroom', '1 bedroom', '2 bedroom', '3 bedroom', '4 bedroom', '5 bedroom']


def table(wb, name):
    for ws in wb.worksheets:
        if name in ws.tables:
            rows = list(ws[ws.tables[name].ref])
            head = [c.value for c in rows[0]]
            return [dict(zip(head, (c.value for c in r))) for r in rows[1:]]
    raise SystemExit(f'table {name} not found: the workbook layout changed; read it by hand')


def main(path: str) -> int:
    data = Path(path).read_bytes()
    wb = openpyxl.load_workbook(path, data_only=True)
    admin = wb['Admin']
    title = wb['Application Instructions']['A1'].value
    basis = {r['Bedrooms']: {'elevator': r['Elevator Base'], 'walk_up': r['Non-Elevator Base']}
             for r in table(wb, 'Basis_Limit_Test_Table')}
    if list(basis) != BEDROOMS:
        raise SystemExit(f'unexpected bedroom rows {list(basis)}')
    zones = {str(r['Zone']): r['Factor'] for r in table(wb, 'Zone_Factor_Table')}
    head = [c.value for c in admin[125][:6]]
    if head[:3] != ['County', 'Zone', 'Low income Score'] or head[5] != 'IS DDA':
        raise SystemExit(f'county table header moved: {head}')
    fips = {f['properties']['NAME']: f['properties']['GEOID'] for f in
            json.loads((ROOT / 'data' / 'co-county-boundaries.json').read_text())['features']}
    counties = {}
    for r in admin.iter_rows(min_row=126, max_row=189, max_col=6, values_only=True):
        if r[0] not in fips:
            raise SystemExit(f'county {r[0]!r} is not in data/co-county-boundaries.json')
        counties[fips[r[0]]] = {'name': r[0], 'zone': r[1], 'lowest_ami_points': r[2], 'chfa_dda': bool(r[5])}
    if len(counties) != 64:
        raise SystemExit(f'{len(counties)} counties, expected 64')
    counties = dict(sorted(counties.items()))
    cities = {r['City']: r['Score'] for r in table(wb, 'City_Table') if r.get('City')}
    out = {
        'schema': 'chfa-common-app-tables/1',
        'status': 'draft',
        'status_note': ('Copied from CHFA\'s 2027 Common Application workbook, which follows the 2027-28 QAP '
                        'Third Draft (posted 2026-09-11). Not adopted: label every figure draft.'),
        'source': {'title': title, 'url': SOURCE_URL, 'page_url': PAGE_URL,
                   'sha256': hashlib.sha256(data).hexdigest(), 'tab': 'Admin (hidden)'},
        'basis_limit_note': admin['C91'].value,
        'basis_limit_per_unit_base': basis,
        'zone_factor': zones,
        'building_type_column': {'Elevator': 'elevator', 'Walk-up': 'walk_up'},
        'method_three': ('limit per unit = TRUNC(zone_factor[county zone] x base[bedrooms][building type]); '
                         'adjusted eligible basis = sum(units x limit per unit); x 1.3 if QCT or DDA; '
                         'x applicable fraction; x applicable percentage (Calculations rows 497-516)'),
        'counties': counties,
        'lowest_ami_points_line': 'Counties with lowest Area Median Income (AMI) (Calculations row 312)',
        'city_cost_burden_points': dict(sorted(cities.items())),
        'city_cost_burden_note': ('Score related to the percentage of renters severely housing cost burdened in city; '
                                  'looked up by exact city name, 0 when the city is not listed '
                                  '(Calculations row 316). ' + str(admin['C223'].value)),
    }
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False) + '\n')
    print(f'wrote {OUT.relative_to(ROOT)}: {len(counties)} counties, {len(cities)} cities')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1]))
